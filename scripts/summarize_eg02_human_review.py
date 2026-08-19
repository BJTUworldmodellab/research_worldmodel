#!/usr/bin/env python3
"""Summarize EG-02 A/D human annotations from XLSX files."""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ALLOWED_LABELS = {"satisfied", "not_satisfied", "uncertain", "not_judgable"}
ANNOTATION_COLUMNS = [
    "annotation_id",
    "scene_id",
    "room_type",
    "layout_variant",
    "relation_id",
    "subject_class",
    "predicate",
    "object_class",
    "image_or_view_path",
    "human_label",
    "missing_object",
    "visibility_quality",
    "annotator_id",
    "notes",
    "baseline_satisfied",
    "repair_satisfied",
]


def normalize_cell(value: object) -> str:
    return "" if value is None else str(value).strip()


def read_xlsx(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"{path} is empty")
    header = [normalize_cell(value) for value in rows[0]]
    index = {name: offset for offset, name in enumerate(header)}
    required = {"annotation_id", "human_label"}
    missing = required.difference(index)
    if missing:
        raise ValueError(f"{path} missing columns: {sorted(missing)}")

    records: list[dict[str, str]] = []
    seen: set[str] = set()
    for raw in rows[1:]:
        record = {name: normalize_cell(raw[offset]) if offset < len(raw) else "" for name, offset in index.items()}
        annotation_id = record["annotation_id"]
        if not annotation_id:
            continue
        if annotation_id in seen:
            raise ValueError(f"{path} duplicate annotation_id: {annotation_id}")
        seen.add(annotation_id)
        label = record["human_label"].lower()
        if label not in ALLOWED_LABELS:
            raise ValueError(f"{path} invalid label for {annotation_id}: {record['human_label']!r}")
        record["human_label"] = label
        records.append(record)
    return records


def read_reference_csv(path: Path | None) -> dict[str, dict[str, str]]:
    if path is None or not path.exists():
        return {}
    with path.open(newline="", encoding="utf-8") as handle:
        return {row["annotation_id"]: row for row in csv.DictReader(handle) if row.get("annotation_id")}


def write_csv(path: Path, rows: list[dict[str, str]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def cohens_kappa(labels_a: dict[str, str], labels_b: dict[str, str]) -> tuple[float, float, float, int]:
    ids_a = set(labels_a)
    ids_b = set(labels_b)
    if ids_a != ids_b:
        raise ValueError(f"annotation id mismatch: only_a={sorted(ids_a - ids_b)[:5]}, only_b={sorted(ids_b - ids_a)[:5]}")
    ids = sorted(ids_a)
    total = len(ids)
    observed = sum(labels_a[item] == labels_b[item] for item in ids) / total
    label_set = sorted(set(labels_a.values()) | set(labels_b.values()))
    counts_a = Counter(labels_a[item] for item in ids)
    counts_b = Counter(labels_b[item] for item in ids)
    expected = sum((counts_a[label] / total) * (counts_b[label] / total) for label in label_set)
    kappa = 1.0 if expected == 1.0 and observed == 1.0 else (observed - expected) / (1.0 - expected)
    return kappa, observed, expected, total


def auto_label(value: str) -> str:
    return "satisfied" if str(value).strip() == "1" else "not_satisfied"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--annotator-a", required=True)
    parser.add_argument("--annotator-d", required=True)
    parser.add_argument("--reference-csv")
    parser.add_argument("--output-dir", default="results/eurographics2027/eg02_human_review")
    args = parser.parse_args()

    output_dir = Path(args.output_dir)
    rows_a = read_xlsx(Path(args.annotator_a))
    rows_d = read_xlsx(Path(args.annotator_d))
    reference = read_reference_csv(Path(args.reference_csv) if args.reference_csv else None)
    for rows in (rows_a, rows_d):
        for row in rows:
            ref = reference.get(row["annotation_id"], {})
            for key in ("image_or_view_path", "source_text", "subject_center_xyz", "object_center_xyz"):
                if not row.get(key) and ref.get(key):
                    row[key] = ref[key]
    write_csv(output_dir / "eg02_human_review_sample_A.csv", rows_a, ANNOTATION_COLUMNS)
    write_csv(output_dir / "eg02_human_review_sample_D.csv", rows_d, ANNOTATION_COLUMNS)

    by_id_a = {row["annotation_id"]: row for row in rows_a}
    by_id_d = {row["annotation_id"]: row for row in rows_d}
    labels_a = {key: row["human_label"] for key, row in by_id_a.items()}
    labels_d = {key: row["human_label"] for key, row in by_id_d.items()}
    kappa_all, observed_all, expected_all, total_all = cohens_kappa(labels_a, labels_d)

    binary_ids = sorted(
        key for key in labels_a
        if labels_a[key] in {"satisfied", "not_satisfied"}
        and labels_d.get(key) in {"satisfied", "not_satisfied"}
    )
    kappa_binary, observed_binary, expected_binary, total_binary = cohens_kappa(
        {key: labels_a[key] for key in binary_ids},
        {key: labels_d[key] for key in binary_ids},
    )

    disagreements = []
    for key in sorted(labels_a):
        if labels_a[key] != labels_d[key]:
            row = by_id_a[key]
            disagreements.append(
                {
                    "annotation_id": key,
                    "room_type": row.get("room_type", ""),
                    "predicate": row.get("predicate", ""),
                    "subject_class": row.get("subject_class", ""),
                    "object_class": row.get("object_class", ""),
                    "label_a": labels_a[key],
                    "label_d": labels_d[key],
                    "image_or_view_path": row.get("image_or_view_path", ""),
                    "source_text": row.get("source_text", ""),
                }
            )

    room_stats: dict[str, Counter] = defaultdict(Counter)
    for key, row in by_id_a.items():
        room = row.get("room_type", "")
        room_stats[room]["total"] += 1
        room_stats[room]["agree"] += int(labels_a[key] == labels_d[key])
        if labels_a[key] in {"satisfied", "not_satisfied"} and labels_d[key] in {"satisfied", "not_satisfied"}:
            room_stats[room]["binary_total"] += 1
            room_stats[room]["binary_agree"] += int(labels_a[key] == labels_d[key])

    machine_rows = []
    for key in sorted(by_id_a):
        row = by_id_a[key]
        if labels_a[key] in {"satisfied", "not_satisfied"}:
            machine_rows.append(("A", key, row["room_type"], labels_a[key], auto_label(row.get("repair_satisfied", "0"))))
        if labels_d[key] in {"satisfied", "not_satisfied"}:
            machine_rows.append(("D", key, row["room_type"], labels_d[key], auto_label(row.get("repair_satisfied", "0"))))
    machine_agree = sum(human == machine for _, _, _, human, machine in machine_rows)

    passed = kappa_all >= 0.70
    summary = {
        "task": "EG-02 human review",
        "status": "PASS_KAPPA_GATE" if passed else "FAIL_KAPPA_GATE",
        "threshold": 0.70,
        "n_relations": total_all,
        "label_counts_a": dict(Counter(labels_a.values())),
        "label_counts_d": dict(Counter(labels_d.values())),
        "all_label_kappa": round(kappa_all, 6),
        "all_label_observed_agreement": round(observed_all, 6),
        "all_label_expected_agreement": round(expected_all, 6),
        "binary_kappa": round(kappa_binary, 6),
        "binary_observed_agreement": round(observed_binary, 6),
        "binary_expected_agreement": round(expected_binary, 6),
        "binary_n_relations": total_binary,
        "disagreement_count": len(disagreements),
        "room_stats": {
            room: {
                "total": counts["total"],
                "agreement": round(counts["agree"] / counts["total"], 6),
                "binary_total": counts["binary_total"],
                "binary_agreement": round(counts["binary_agree"] / counts["binary_total"], 6) if counts["binary_total"] else None,
            }
            for room, counts in sorted(room_stats.items())
        },
        "repair_auto_vs_human_binary": {
            "n_annotation_decisions": len(machine_rows),
            "agreement": round(machine_agree / len(machine_rows), 6) if machine_rows else None,
        },
    }

    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "summary.json").write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    write_csv(
        output_dir / "disagreements.csv",
        disagreements,
        ["annotation_id", "room_type", "predicate", "subject_class", "object_class", "label_a", "label_d", "image_or_view_path", "source_text"],
    )
    adjudication_rows = [
        {
            **row,
            "final_label": "",
            "adjudicator_id": "",
            "adjudication_notes": "",
        }
        for row in disagreements
    ]
    write_csv(
        output_dir / "adjudication_template.csv",
        adjudication_rows,
        [
            "annotation_id",
            "room_type",
            "predicate",
            "subject_class",
            "object_class",
            "label_a",
            "label_d",
            "final_label",
            "adjudicator_id",
            "adjudication_notes",
            "image_or_view_path",
            "source_text",
        ],
    )
    summary_rows = [
        {"metric": key, "value": value}
        for key, value in summary.items()
        if not isinstance(value, (dict, list))
    ]
    write_csv(output_dir / "summary.csv", summary_rows, ["metric", "value"])

    print(f"n_relations={total_all}")
    print(f"all_label_kappa={kappa_all:.4f}")
    print(f"binary_n_relations={total_binary}")
    print(f"binary_kappa={kappa_binary:.4f}")
    print(f"disagreement_count={len(disagreements)}")
    print(f"output_dir={output_dir}")
    if not passed:
        print("EG-02 HUMAN REVIEW: FAIL all-label kappa below 0.70")
        return 1
    print("EG-02 HUMAN REVIEW: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
